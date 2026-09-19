"""services/doc_quiz — question types, validator, dedup, selection, calibration, DOCX, end-to-end with a stub LLM."""

import asyncio
import io

import pytest

from services.doc_quiz import calibration as calib
from services.doc_quiz import generate as gen
from services.doc_quiz import items as it

CHUNK = ("The Consumer Price Index (CPI) compiled by the NSO uses the Laspeyres formula with base year 2012. "
         "In a sample survey, stratified sampling divides the population into homogeneous strata before "
         "drawing a simple random sample from each stratum. If 25 of 200 households report a purchase, "
         "the share is 12.5 percent.")
DOC = f"--- Page 1 ---\nIntro text.\n\n--- Page 2 ---\n{CHUNK}"


def _chunks():
    cs = gen.locate_chunks(DOC, [CHUNK])
    return {c.id: c for c in cs}, {c.id for c in cs}


def _mcq(**kw):
    base = {"type": "mcq", "chunk": "c1", "quote": "the NSO uses the Laspeyres formula with base year 2012",
            "question": "Which index formula is used to compile the CPI in India?",
            "options": ["Laspeyres", "Paasche", "Fisher ideal", "Törnqvist"], "correct_answer": 0,
            "explanation": "The CPI uses the Laspeyres base-weighted formula.",
            "why_wrong": ["", "Paasche uses current weights.", "Fisher is a geometric mean of both.", "Chain index."],
            "difficulty": "Easy"}
    base.update(kw)
    return base


# ── items ─────────────────────────────────────────────────────────────────────

def test_is_correct_by_type():
    assert it.is_correct({"type": "mcq", "options": list("abcd"), "correct_answer": 2}, 2)
    assert not it.is_correct({"type": "mcq", "options": list("abcd"), "correct_answer": 2}, "x")
    assert it.is_correct({"type": "true_false", "options": ["True", "False"], "correct_answer": 1}, 1)
    ms = {"type": "multi_select", "options": list("abcde"), "correct_answers": [0, 3]}
    assert it.is_correct(ms, [3, 0]) and not it.is_correct(ms, [0]) and not it.is_correct(ms, [0, 3, 4])
    fb = {"type": "fill_blank", "answer_text": "stratified sampling", "accepted_answers": ["stratification"]}
    assert it.is_correct(fb, "Stratified  Sampling.") and it.is_correct(fb, "stratification")
    assert it.is_correct(fb, "stratifed sampling")            # one-letter typo tolerated
    assert not it.is_correct(fb, "cluster sampling")
    num = {"type": "numeric", "numeric_answer": 12.5}
    assert it.is_correct(num, "12.5%") and it.is_correct(num, 12.45) and not it.is_correct(num, 13)
    assert it.is_correct(num, "१२.५")                          # Devanagari digits


def test_media_style_questions_default_to_mcq():
    class Q:  # a media QuizQuestion copy carries no type
        question, options, correct_answer, explanation = "q?", ["a", "b", "c", "d"], 1, ""
    assert it.qtype(Q()) == "mcq" and it.is_correct(Q(), 1)


def test_numeric_feedback_diagnoses_percentage_and_sign():
    q = {"type": "numeric", "numeric_answer": 12.5, "unit": "%", "solution": "25/200*100"}
    assert "fraction" in it.feedback(q, "0.125")
    assert "sign" in it.feedback(q, "-12.5")
    assert "Close" in it.feedback(q, "12.0")
    assert "Working: 25/200*100" in it.feedback(q, "40")
    assert it.feedback(q, "12.5") == ""


def test_choice_feedback_uses_option_rationale():
    q = {"type": "mcq", "options": ["Laspeyres", "Paasche", "Fisher", "Chain"], "correct_answer": 0,
         "why_wrong": ["", "Paasche uses current-period weights.", "", ""]}
    fb = it.feedback(q, 1)
    assert "Paasche" in fb and "current-period weights" in fb
    ms = {"type": "multi_select", "options": list("abcd"), "correct_answers": [0, 1], "why_wrong": ["", "", "c is x", ""]}
    fb = it.feedback(ms, [0, 2])
    assert "should not be selected" in fb and "missed" in fb


def test_safe_eval():
    assert it.safe_eval("(25/200)*100") == 12.5
    assert it.safe_eval("sqrt(16) + 2^3") == 12
    assert abs(it.safe_eval("mean(2, 4, 6)") - 4) < 1e-9
    for bad in ("__import__('os')", "open('x')", "2**999", "a+1"):
        with pytest.raises((ValueError, SyntaxError)):
            it.safe_eval(bad)


def test_item_key_stable_across_whitespace_and_case():
    a = {"type": "mcq", "question": "What is CPI?", "options": ["x", "y"], "correct_answer": 0}
    b = {"type": "mcq", "question": "what  is CPI ?", "options": ["X", "z"], "correct_answer": 0}
    assert it.item_key(a) == it.item_key(b)


# ── validator ─────────────────────────────────────────────────────────────────

def test_locator_from_page_markers():
    cmap, _ = _chunks()
    assert cmap["c1"].locator == "Page 2"
    assert "--- Page" not in cmap["c1"].text


def test_valid_mcq_is_cited():
    cmap, allowed = _chunks()
    q = gen.validate(_mcq(), cmap, allowed, {}, ["mcq"])
    assert q and q["citations"][0]["locator"] == "Page 2"
    assert "Laspeyres" in q["citations"][0]["passage"]
    assert q["why_wrong"][0] == ""


@pytest.mark.parametrize("patch,reason", [
    ({"quote": "The CPI is published by the Reserve Bank every week"}, "quote_not_in_source"),
    ({"chunk": "c9"}, "cited_unknown_chunk"),
    ({"options": ["Laspeyres", "Paasche", "Fisher ideal", "None of the above"]}, "catch_all_option"),
    ({"options": ["Laspeyres", "Laspeyre", "Fisher ideal", "Paasche"]}, "near_duplicate_options"),
    ({"question": "Which formula does the author mention on page 2?"}, "about_the_document_not_the_subject"),
    ({"options": ["Base year 2015", "Base year 2004", "Base year 2001", "Base year 1999"],
      "question": "What is the base year of the CPI series?"}, "answer_number_not_in_source"),
    ({"options": ["Laspeyres base-weighted formula applied to all item groups every month", "Paasche",
                  "Fisher", "Chain"]}, "answer_length_giveaway"),
    ({"options": ["2012", "Paasche", "Fisher", "Chain"],
      "question": "Which base year is used to compile the CPI?"}, "implausible_distractor_type"),
])
def test_validator_rejections(patch, reason):
    cmap, allowed = _chunks()
    stats = {}
    assert gen.validate(_mcq(**patch), cmap, allowed, stats, ["mcq"]) is None
    assert stats == {reason: 1}


def test_numeric_expression_checked():
    cmap, allowed = _chunks()
    base = {"type": "numeric", "chunk": "c1", "quote": "If 25 of 200 households report a purchase",
            "question": "What percentage of households reported a purchase?", "explanation": "25/200 × 100.",
            "numeric_answer": 12.5, "unit": "%", "expression": "25/200*100", "solution": "25 ÷ 200 × 100 = 12.5"}
    ok = gen.validate(base, cmap, allowed, {}, ["numeric"])
    assert ok and ok["numeric_answer"] == 12.5 and ok["answer_type"] == "number"
    stats = {}
    assert gen.validate({**base, "numeric_answer": 15}, cmap, allowed, stats, ["numeric"]) is None
    assert stats == {"expression_mismatch": 1}
    stats = {}
    assert gen.validate({**base, "expression": "35/280*100"}, cmap, allowed, stats, ["numeric"]) is None
    assert stats == {"expression_inputs_not_in_source": 1}


def test_fill_blank_and_true_false_rules():
    cmap, allowed = _chunks()
    fb = {"type": "fill_blank", "chunk": "c1", "quote": "stratified sampling divides the population into homogeneous strata",
          "question": "In a sample survey, _____ divides the population into homogeneous strata.",
          "answer_text": "stratified sampling", "explanation": "Stated in the source."}
    assert gen.validate(fb, cmap, allowed, {}, ["fill_blank"])["answer_text"] == "stratified sampling"
    stats = {}
    assert gen.validate({**fb, "question": "Stratified sampling: _____ divides the population."},
                        cmap, allowed, stats, ["fill_blank"]) is None
    assert stats == {"fill_blank_answer_in_stem": 1}
    tf = {"type": "true_false", "chunk": "c1", "quote": "the NSO uses the Laspeyres formula with base year 2012",
          "question": "The CPI compiled by the NSO uses the Paasche formula.", "correct_answer": False,
          "explanation": "It uses the Laspeyres formula, not Paasche."}
    q = gen.validate(tf, cmap, allowed, {}, ["true_false"])
    assert q["options"] == ["True", "False"] and q["correct_answer"] == 1 and q["why_wrong"][0]
    stats = {}
    assert gen.validate({**tf, "explanation": ""}, cmap, allowed, stats, ["true_false"]) is None
    assert stats == {"false_statement_without_correction": 1}


def test_lexical_dedup_when_no_embedder(monkeypatch):
    import ai.embedder
    monkeypatch.setattr(ai.embedder, "get_embedder", lambda *_a, **_k: (_ for _ in ()).throw(RuntimeError("offline")))
    cmap, allowed = _chunks()
    a = gen.validate(_mcq(), cmap, allowed, {}, ["mcq"])
    b = gen.validate(_mcq(question="Which index formula is used when compiling the CPI in India?"), cmap, allowed, {}, ["mcq"])
    stats = {}
    kept, method = gen.plausibility_and_dedup([a, b], stats)
    assert method == "lexical" and len(kept) == 1 and stats == {"duplicate": 1}


def test_selection_honours_type_quotas():
    def cand(t, d, c):
        return {"type": t, "difficulty": d, "citations": [{"chunkId": c}]}
    pool = [cand("mcq", "Medium", "c1")] * 6 + [cand("numeric", "Hard", "c2"), cand("true_false", "Easy", "c3")]
    out = gen.select(pool, 4, ["mcq", "numeric", "true_false"], "Medium")
    assert sorted(q["type"] for q in out) == ["mcq", "mcq", "numeric", "true_false"]
    all_types = ["mcq", "true_false", "multi_select", "fill_blank", "numeric"]
    assert set(gen.type_quotas(all_types, 5).values()) == {1}          # one of each
    assert gen.type_quotas(all_types, 12)["mcq"] == 4                   # MCQ gets a double share


# ── calibration ───────────────────────────────────────────────────────────────

def test_calibration_prior_then_data():
    assert calib.calibrate("Hard")["difficulty"] == "Hard"
    assert calib.calibrate("Hard", 3, 3, 3 * 2.0)["source"] == "llm_tag"          # too few responses
    # Tagged Hard, but weak learners (θ≈1.8) all get it right → it is Easy.
    c = calib.calibrate("Hard", 20, 19, 20 * 1.8)
    assert c["source"] == "response_data" and c["difficulty"] == "Easy" and not c["agreesWithLlm"]
    # Tagged Easy, but strong learners (θ≈3.5) mostly miss it → Hard.
    assert calib.calibrate("Easy", 20, 4, 20 * 3.5)["difficulty"] == "Hard"
    # Consistent data keeps the tag.
    assert calib.calibrate("Medium", 20, 12, 20 * 2.7)["difficulty"] == "Medium"


# ── DOCX ──────────────────────────────────────────────────────────────────────

def test_extract_docx_keeps_order_headings_and_tables():
    import docx
    from services.doc_quiz.extract import extract_docx

    d = docx.Document()
    d.add_heading("Price collection", level=1)
    d.add_paragraph("Prices are collected monthly from selected markets.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Weight"
    t.cell(1, 0).text, t.cell(1, 1).text = "Cereals", "9.67"
    buf = io.BytesIO()
    d.save(buf)
    text, sections = extract_docx(buf.getvalue())
    assert sections == 1
    assert text.index("--- Section 1: Price collection ---") < text.index("collected monthly") < text.index("Cereals | 9.67")


# ── End-to-end with a stub LLM ────────────────────────────────────────────────

def test_generate_end_to_end(monkeypatch):
    import ai.embedder
    monkeypatch.setattr(ai.embedder, "get_embedder", lambda *_a, **_k: (_ for _ in ()).throw(RuntimeError("offline")))
    canned = {"questions": [
        _mcq(),
        _mcq(question="Which formula for CPI?", quote="invented text never in the chunk at all here"),   # rejected
        {"type": "numeric", "chunk": "c1", "quote": "If 25 of 200 households report a purchase",
         "question": "What percentage of the 200 households reported a purchase?", "explanation": "25/200×100",
         "numeric_answer": 12.5, "unit": "%", "expression": "25/200*100", "difficulty": "Medium"},
        {"type": "true_false", "chunk": "c1", "quote": "stratified sampling divides the population into homogeneous strata",
         "question": "Stratified sampling divides the population into homogeneous strata.", "correct_answer": True,
         "explanation": "As stated.", "difficulty": "Easy"},
    ]}

    async def fake_llm(prompt, *a, **k):
        assert "VERBATIM" in prompt and "why_wrong" in prompt
        return canned
    monkeypatch.setattr(gen, "gemini_json", fake_llm)

    async def fake_tr(texts, lang):
        return [f"[hi] {t}" for t in texts], {"target": lang, "status": "ok"}
    monkeypatch.setattr(gen, "translate_texts", fake_tr)

    qs, report = asyncio.run(gen.generate(DOC, [CHUNK * 2], "Medium", 3,
                                          ["mcq", "numeric", "true_false"], "bi"))
    assert report["accepted"] == 3 and report["rejected"] == {"quote_not_in_source": 1}
    assert {q["type"] for q in qs} == {"mcq", "numeric", "true_false"}
    assert all(q["citations"][0]["quote"] for q in qs)
    tf = next(q for q in qs if q["type"] == "true_false")
    assert tf["translations"]["hi"]["options"] == ["सही", "गलत"]
    assert next(q for q in qs if q["type"] == "mcq")["translations"]["hi"]["question"].startswith("[hi]")


def test_quote_locator_inside_multi_section_chunk():
    text = "--- Section 1: CPI ---\nThe CPI uses the Laspeyres formula.\n\n--- Section 2: Sampling ---\nStrata are homogeneous groups of units."
    (c,) = gen.locate_chunks(text, [text])
    assert c.locator == "Sections 1–2"
    assert c.locate("Strata are homogeneous groups of units.") == "Section 2: Sampling"
    assert c.locate("The CPI uses the Laspeyres formula.") == "Section 1: CPI"


def test_courses_for_topics_prefers_competency_courses(monkeypatch):
    import numpy as np
    import ai.embedder
    from services import practice_assessment as pa

    class Doc:
        def __init__(self, i):
            self.identifier, self.name, self.comp_levels, self.duration_hrs, self.rating = f"do_{i}", f"Course {i}", {"C1": 2}, 3.0, 4.5

    class Eng:
        _catalog = [Doc(0), Doc(1), Doc(2)]
        _embeddings = np.eye(3, dtype="float32")
        _comp_index = {"C1": [1, 2]}

    class Emb:
        def encode(self, texts, **_):
            return np.array([[1.0, 0.0, 0.0], [0.0, 0.2, 0.9]], dtype="float32")[:len(texts)]
    monkeypatch.setattr(ai.embedder, "get_embedder", lambda *_a, **_k: Emb())
    out = pa.courses_for_topics(Eng(), ["a", "b"], "C1")
    # Course 0 matches "a" best overall but is not tagged C1 → the tagged pool wins.
    assert [c["courseId"] for c in out] == ["do_1", "do_2"] and out[0]["fromCompetency"]
    assert pa.courses_for_topics(None, ["a"], "C1") == [None]
